"""Tests for Mobile Capture XLSForm Generator, Sync Engine, and dMRV ACVA CLI."""
from __future__ import annotations

import datetime
import io
from pathlib import Path
import openpyxl
import pytest

from cli.main import main
from dmrv_core.data.mobile_forms import (
    MobileSyncEngine,
    MobileSyncRecord,
    XLSFormGenerator,
)
from reports.bundle import ReproducibilityBundle


def test_xlsform_generation() -> None:
    xlsx_bytes = XLSFormGenerator.generate_forestry_xlsform(
        species_list=["Tectona grandis", "Dalbergia sissoo"],
        form_id="test_forestry_v1",
        form_title="Test Forestry Form",
    )
    assert len(xlsx_bytes) > 0

    # Load back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    assert "survey" in sheet_names
    assert "choices" in sheet_names
    assert "settings" in sheet_names

    ws_survey = wb["survey"]
    headers = [cell.value for cell in ws_survey[1]]
    assert "type" in headers
    assert "name" in headers
    assert "constraint" in headers

    # Check constraint on DBH
    found_dbh = False
    for row in ws_survey.iter_rows(values_only=True):
        if row[1] == "dbh_cm":
            assert ". >= 5.0 and . <= 300.0" in str(row[4])
            found_dbh = True
    assert found_dbh is True


def test_mobile_sync_validation_and_rejection() -> None:
    now = datetime.datetime.now(datetime.timezone.utc)

    # 1. Valid record
    rec_valid = MobileSyncRecord(
        client_uuid="uuid-001",
        device_id="DEV_SAMSUNG_A12",
        operator_id="OP_RAVI",
        timestamp_utc=now,
        gps_lat=24.5,
        gps_lon=78.2,
        gps_accuracy_m=4.2,
        form_type="trees",
        data={"plot_id": "P01", "tree_tag": "T101", "dbh_cm": 28.5, "height_m": 14.2},
    )

    # 2. Record with bad GPS (> 30 m)
    rec_bad_gps = MobileSyncRecord(
        client_uuid="uuid-002",
        device_id="DEV_SAMSUNG_A12",
        operator_id="OP_RAVI",
        timestamp_utc=now,
        gps_lat=24.5,
        gps_lon=78.2,
        gps_accuracy_m=45.0,  # Exceeds 30 m threshold
        form_type="trees",
        data={"plot_id": "P01", "tree_tag": "T102", "dbh_cm": 20.0, "height_m": 12.0},
    )

    # 3. Duplicate tree tag in same plot
    rec_dup_tag = MobileSyncRecord(
        client_uuid="uuid-003",
        device_id="DEV_SAMSUNG_A12",
        operator_id="OP_RAVI",
        timestamp_utc=now,
        gps_lat=24.5,
        gps_lon=78.2,
        gps_accuracy_m=3.5,
        form_type="trees",
        data={"plot_id": "P01", "tree_tag": "T101", "dbh_cm": 30.0, "height_m": 15.0},
    )

    batch_res = MobileSyncEngine.process_sync_batch([rec_valid, rec_bad_gps, rec_dup_tag])
    assert batch_res.total_records == 3
    assert batch_res.accepted_records == 1
    assert batch_res.rejected_records == 2
    assert batch_res.sync_status == "PARTIAL"


def test_cli_bundle_verification(tmp_path: Path) -> None:
    from dmrv_core.equations import CalculationContext
    from methodologies.plugins.fr05_002 import BMFR05002Methodology

    plugin = BMFR05002Methodology()
    ctx = CalculationContext(project_id="PRJ_CLI_TEST")

    project_meta = {
        "project_id": "PRJ_CLI_TEST",
        "start_year": 2026,
    }
    period_data = {
        "start_date": "2026-01-01",
        "end_date": "2026-12-31",
        "ar0004_project": {
            "strata": [{"name": "S1", "area_ha": 50.0, "plot_biomass_list": [40.0, 45.0, 42.0]}],
        },
        "ar0004_baseline": {
            "strata": [{"name": "S1", "area_ha": 50.0, "plot_biomass_list": [0.0, 0.0, 0.0]}],
        },
        "ar0003": {
            "C_DW_tCO2e": 10.0,
            "C_LI_tCO2e": 5.0,
        },
        "ar0006": {
            "SOC_REF": 47.0,
            "f_LU": 0.82,
            "f_MG": 1.0,
            "f_IN": 0.92,
            "t_PREP": 1,
            "A_i_ha": 50.0,
            "disturbance_gt_10pct": False,
        },
        "ar0002": {
            "accumulated_burnt_area_ha": 0.0,
            "project_area_ha": 50.0,
        },
        "ar0005": {
            "zero_leakage_grazing_case": "a",
        },
        "C_C_tCO2e": 0.0,
        "has_prior_crediting": False,
    }
    res = plugin.ex_post(ctx, project_meta, period_data)

    bundle_zip_path = tmp_path / "valid_bundle.zip"
    ReproducibilityBundle.create_bundle(
        zip_path=bundle_zip_path,
        project_metadata=project_meta,
        inputs=period_data,
        methodology_result=res,
        engine_version="0.1.0",
    )
    assert bundle_zip_path.exists()

    # Run CLI command: dmrv run --bundle <zip>
    exit_code = main(["run", "--bundle", str(bundle_zip_path)])
    assert exit_code == 0

    # Non-existent bundle returns exit code 1
    exit_code_missing = main(["run", "--bundle", str(tmp_path / "non_existent.zip")])
    assert exit_code_missing == 1
